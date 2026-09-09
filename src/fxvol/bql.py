from __future__ import annotations
import math
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from .constants import SECURITY_CODES
TABLE_NS = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
TEXT_NS = 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'
OFFICE_NS = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
SECURITY_PATTERN = re.compile('^(?P<pair>[A-Z]{6})(?P<code>V|25R|25B|10R|10B)(?P<tenor>1W|1M|3M|6M|1Y) BGN Curncy$')
FIELD_ALIASES = {'#last': 'last', '#bid': 'bid', '#ask': 'ask', 'px_last': 'last', 'px_bid': 'bid', 'px_ask': 'ask'}

def _normalise_field(value: str) -> str | None:
    field = value.strip().lower()
    if '(' in field:
        field = field.split('(', 1)[0]
    return FIELD_ALIASES.get(field)

@dataclass(frozen=True)
class PanelRecord:
    date: date
    pair: str
    tenor: str
    quote_type: str
    last: float | None
    bid: float | None
    ask: float | None

    @property
    def spread(self) -> float | None:
        if self.bid is None or self.ask is None:
            return None
        return self.ask - self.bid

    @property
    def is_observed(self) -> bool:
        return self.last is not None

    @property
    def midpoint_difference(self) -> float | None:
        if self.last is None or self.bid is None or self.ask is None:
            return None
        return abs(self.last - (self.bid + self.ask) / 2)

def _read_nonempty_rows(path: Path) -> list[list[str]]:
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read('content.xml'))
    table = root.find(f'.//{{{TABLE_NS}}}table')
    if table is None:
        raise ValueError(f'No worksheet found in {path}')
    rows: list[list[str]] = []
    for row in table.findall(f'{{{TABLE_NS}}}table-row'):
        values: list[str] = []
        for cell in row.findall(f'{{{TABLE_NS}}}table-cell'):
            repeat = int(cell.attrib.get(f'{{{TABLE_NS}}}number-columns-repeated', '1'))
            text = ' '.join((''.join(paragraph.itertext()) for paragraph in cell.findall(f'{{{TEXT_NS}}}p'))).strip()
            raw = cell.attrib.get(f'{{{OFFICE_NS}}}value') or cell.attrib.get(f'{{{OFFICE_NS}}}date-value')
            values.extend([raw if raw is not None else text] * min(repeat, 1000))
        if any((str(value).strip() for value in values)):
            rows.append(values)
    return rows

def _read_xlsx_nonempty_rows(path: Path) -> list[list[str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            values = ['' if value is None else str(value) for value in row]
            while values and values[-1] == '':
                values.pop()
            if any((value.strip() for value in values)):
                rows.append(values)
        return rows
    finally:
        workbook.close()

def read_nonempty_bql_rows(path: str | Path) -> list[list[str]]:
    path = Path(path)
    if path.suffix.lower() == '.ods':
        return _read_nonempty_rows(path)
    if path.suffix.lower() in {'.xlsx', '.xlsm'}:
        return _read_xlsx_nonempty_rows(path)
    raise ValueError(f'Unsupported BQL workbook extension: {path}')

def _number(value: str) -> float | None:
    if value in ('', '#N/A'):
        return None
    number = float(value)
    return number if math.isfinite(number) else None

def _date(value: str) -> date:
    try:
        return datetime.fromisoformat(value).date()
    except ValueError:
        number = float(value)
        return from_excel(number).date()

def read_bql_workbook(path: str | Path, weekdays_only: bool=True) -> list[PanelRecord]:
    rows = read_nonempty_bql_rows(Path(path))
    if len(rows) < 3 or rows[1][0] != 'DATES':
        raise ValueError(f'Unexpected BQL workbook structure in {path}')
    securities = rows[0]
    fields = rows[1]
    column_map: dict[tuple[str, str, str, str], dict[str, int]] = {}
    for index in range(1, min(len(securities), len(fields))):
        match = SECURITY_PATTERN.match(securities[index])
        field = _normalise_field(fields[index])
        if not match or field is None:
            continue
        key = (match.group('pair'), match.group('tenor'), SECURITY_CODES[match.group('code')], securities[index])
        column_map.setdefault(key, {})[field] = index
    records: list[PanelRecord] = []
    for row in rows[2:]:
        try:
            observation_date = _date(row[0])
        except (ValueError, IndexError):
            continue
        if weekdays_only and observation_date.weekday() >= 5:
            continue
        for (pair, tenor, quote_type, _), indexes in column_map.items():
            values = {field: _number(row[index]) if index < len(row) else None for field, index in indexes.items()}
            records.append(PanelRecord(date=observation_date, pair=pair, tenor=tenor, quote_type=quote_type, last=values.get('last'), bid=values.get('bid'), ask=values.get('ask')))
    return records

def read_bql_ods(path: str | Path, weekdays_only: bool=True) -> list[PanelRecord]:
    return read_bql_workbook(path, weekdays_only=weekdays_only)
